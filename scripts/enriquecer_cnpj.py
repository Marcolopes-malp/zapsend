#!/usr/bin/env python3
"""
Enriquece uma planilha de leads com telefone/e-mail a partir do CNPJ (BrasilAPI).

Uso:
    python3 scripts/enriquecer_cnpj.py leads.csv
    python3 scripts/enriquecer_cnpj.py leads.xlsx --coluna "CNPJ da Empresa"

Gera <arquivo>_com_telefone.csv/.xlsx com as colunas novas no final.
As consultas ficam salvas em <arquivo>.cache.json: se o script parar no meio,
é só rodar de novo que ele continua de onde parou.
"""

import argparse
import csv
import json
import re
import sys
import time
import urllib.error
import urllib.request
from pathlib import Path

API_URL = "https://brasilapi.com.br/api/cnpj/v1/{}"

NOVAS_COLUNAS = [
    "whatsapp_provavel",
    "telefone_1",
    "telefone_2",
    "email",
    "razao_social",
    "nome_fantasia",
    "situacao",
    "municipio",
    "uf",
    "status_consulta",
]


# ---------- CNPJ ----------

def limpar_cnpj(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float):  # Excel guarda CNPJ como número e perde zeros à esquerda
        valor = int(valor)
    digitos = re.sub(r"\D", "", str(valor))
    return digitos.zfill(14) if 0 < len(digitos) <= 14 else digitos


def cnpj_valido(cnpj: str) -> bool:
    if len(cnpj) != 14 or cnpj == cnpj[0] * 14:
        return False

    def digito(base: str, pesos: list[int]) -> str:
        resto = sum(int(d) * p for d, p in zip(base, pesos)) % 11
        return "0" if resto < 2 else str(11 - resto)

    pesos1 = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
    d1 = digito(cnpj[:12], pesos1)
    d2 = digito(cnpj[:12] + d1, [6] + pesos1)
    return cnpj[12:] == d1 + d2


# ---------- Telefone ----------

def normalizar_telefone(bruto: str) -> str:
    """DDD + número, só dígitos, no padrão atual. Vazio se for lixo."""
    # Zeros à esquerda: "011..." (prefixo de operadora) ou "000000000000" (placeholder da Receita)
    d = re.sub(r"\D", "", bruto or "").lstrip("0")
    if len(d) == 10 and d[2] in "6789":
        # Celular no formato antigo, de antes do nono dígito (a Receita guarda muitos assim)
        d = d[:2] + "9" + d[2:]
    return d if len(d) in (10, 11) else ""


def formatar_telefone(bruto: str) -> str:
    d = normalizar_telefone(bruto)
    if len(d) == 11:
        return f"({d[:2]}) {d[2:7]}-{d[7:]}"
    if len(d) == 10:
        return f"({d[:2]}) {d[2:6]}-{d[6:]}"
    return ""


def eh_celular(bruto: str) -> bool:
    # Celular no Brasil: DDD + 9 dígitos começando com 9 (fixos começam com 2 a 5)
    d = normalizar_telefone(bruto)
    return len(d) == 11 and d[2] == "9"


# ---------- BrasilAPI ----------

def consultar(cnpj: str, tentativas: int = 5) -> dict:
    req = urllib.request.Request(
        API_URL.format(cnpj),
        headers={"User-Agent": "LeadZap-Enriquecedor/1.0", "Accept": "application/json"},
    )
    espera = 5
    for _ in range(tentativas):
        try:
            with urllib.request.urlopen(req, timeout=30) as resp:
                return {"ok": True, "dados": json.load(resp)}
        except urllib.error.HTTPError as e:
            if e.code in (400, 404):
                return {"ok": False, "erro": "CNPJ não encontrado"}
            if e.code == 429 or e.code >= 500:  # limite de requisições / instabilidade
                print(f"    API respondeu {e.code}, aguardando {espera}s...")
                time.sleep(espera)
                espera = min(espera * 2, 60)
                continue
            return {"ok": False, "erro": f"HTTP {e.code}"}
        except (urllib.error.URLError, TimeoutError) as e:
            print(f"    Falha de conexão ({e}), aguardando {espera}s...")
            time.sleep(espera)
            espera = min(espera * 2, 60)
    return {"ok": False, "erro": "falhou após várias tentativas", "temporario": True}


def montar_colunas(resultado: dict) -> dict:
    if not resultado.get("ok"):
        return {"status_consulta": resultado.get("erro", "erro")}

    d = resultado["dados"]
    tel1 = formatar_telefone(d.get("ddd_telefone_1"))
    tel2 = formatar_telefone(d.get("ddd_telefone_2"))
    celular = next((t for t in (tel1, tel2) if eh_celular(t)), "")

    return {
        "whatsapp_provavel": celular,
        "telefone_1": tel1,
        "telefone_2": tel2,
        "email": (d.get("email") or "").lower(),
        "razao_social": d.get("razao_social") or "",
        "nome_fantasia": d.get("nome_fantasia") or "",
        "situacao": d.get("descricao_situacao_cadastral") or "",
        "municipio": d.get("municipio") or "",
        "uf": d.get("uf") or "",
        "status_consulta": "ok" if (tel1 or tel2) else "sem telefone cadastrado",
    }


# ---------- Leitura / escrita de planilhas ----------

def ler_csv(caminho: Path):
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            texto = caminho.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    try:
        delimitador = csv.Sniffer().sniff(texto[:5000], delimiters=",;\t").delimiter
    except csv.Error:
        delimitador = ";"
    linhas = list(csv.reader(texto.splitlines(), delimiter=delimitador))
    return linhas[0], linhas[1:], delimitador


def escrever_csv(caminho: Path, cabecalho, linhas, delimitador):
    # utf-8-sig para o Excel abrir os acentos corretamente
    with caminho.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=delimitador)
        w.writerow(cabecalho)
        w.writerows(linhas)


def importar_openpyxl():
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        sys.exit("Para ler .xlsx instale o openpyxl:  pip3 install openpyxl\n"
                 "(ou salve a planilha como CSV e rode de novo)")


def achar_coluna(cabecalho, nome):
    """Índice da coluna do CNPJ, ou None se a aba/arquivo não tiver uma."""
    if nome:
        return cabecalho.index(nome) if nome in cabecalho else None
    for i, c in enumerate(cabecalho):
        if "cnpj" in c.lower():
            return i
    return None


def info_da_linha(valor_cnpj, cache) -> dict:
    cnpj = limpar_cnpj(valor_cnpj)
    if not cnpj_valido(cnpj):
        return {"status_consulta": "CNPJ inválido"}
    if cnpj in cache:
        return montar_colunas(cache[cnpj])
    return {"status_consulta": "não consultado"}


# ---------- Fontes de dados ----------
# CSV: uma tabela só. XLSX: todas as abas que têm coluna de CNPJ, mantendo as demais abas intactas.

class FonteCsv:
    def __init__(self, caminho: Path, coluna):
        self.cabecalho, self.linhas, self.delimitador = ler_csv(caminho)
        self.col = achar_coluna(self.cabecalho, coluna)
        if self.col is None:
            sys.exit(f"Não achei a coluna do CNPJ. Use --coluna. Colunas: {self.cabecalho}")

    def cnpjs(self):
        return [l[self.col] for l in self.linhas if self.col < len(l)]

    def salvar(self, saida: Path, cache):
        novas = []
        for linha in self.linhas:
            info = info_da_linha(linha[self.col] if self.col < len(linha) else "", cache)
            linha = list(linha) + [""] * (len(self.cabecalho) - len(linha))
            novas.append(linha + [info.get(c, "") for c in NOVAS_COLUNAS])
        escrever_csv(saida, self.cabecalho + NOVAS_COLUNAS, novas, self.delimitador)
        return [info_da_linha(v, cache) for v in self.cnpjs()]


class FonteXlsx:
    def __init__(self, caminho: Path, coluna):
        openpyxl = importar_openpyxl()
        self.wb = openpyxl.load_workbook(caminho)  # sem data_only: preserva fórmulas das abas de resumo
        self.abas = []  # (worksheet, índice 1-based da coluna CNPJ)
        for ws in self.wb.worksheets:
            cabecalho = [str(c.value) if c.value is not None else "" for c in ws[1]]
            col = achar_coluna(cabecalho, coluna)
            if col is not None:
                self.abas.append((ws, col + 1))
        if not self.abas:
            sys.exit("Nenhuma aba tem coluna de CNPJ. Use --coluna com o nome exato.")
        print("Abas com CNPJ: " + ", ".join(ws.title for ws, _ in self.abas))

    def cnpjs(self):
        return [
            ws.cell(row=r, column=col).value
            for ws, col in self.abas
            for r in range(2, ws.max_row + 1)
        ]

    def salvar(self, saida: Path, cache):
        infos = []
        for ws, col in self.abas:
            inicio = ws.max_column + 1
            for j, nome in enumerate(NOVAS_COLUNAS):
                ws.cell(row=1, column=inicio + j, value=nome)
            for r in range(2, ws.max_row + 1):
                valor = ws.cell(row=r, column=col).value
                if valor is None:
                    continue
                info = info_da_linha(valor, cache)
                infos.append(info)
                for j, nome in enumerate(NOVAS_COLUNAS):
                    ws.cell(row=r, column=inicio + j, value=info.get(nome, ""))
        self.wb.save(saida)
        return infos


# ---------- Main ----------

def main():
    parser = argparse.ArgumentParser(description="Busca telefone dos leads pelo CNPJ (BrasilAPI)")
    parser.add_argument("arquivo", help="Planilha de leads (.csv ou .xlsx)")
    parser.add_argument("--coluna", help="Nome da coluna do CNPJ (padrão: detecta sozinho)")
    parser.add_argument("--intervalo", type=float, default=1.0,
                        help="Segundos entre consultas, para não estourar o limite da API (padrão: 1)")
    args = parser.parse_args()

    entrada = Path(args.arquivo)
    if not entrada.exists():
        sys.exit(f"Arquivo não encontrado: {entrada}")

    eh_xlsx = entrada.suffix.lower() in (".xlsx", ".xlsm")
    fonte = FonteXlsx(entrada, args.coluna) if eh_xlsx else FonteCsv(entrada, args.coluna)
    saida = entrada.with_name(f"{entrada.stem}_com_telefone{entrada.suffix}")
    cache_path = entrada.with_name(f"{entrada.stem}.cache.json")
    cache = json.loads(cache_path.read_text()) if cache_path.exists() else {}

    valores = [v for v in fonte.cnpjs() if v is not None and str(v).strip()]
    unicos = {limpar_cnpj(v) for v in valores}
    pendentes = sorted(c for c in unicos if cnpj_valido(c) and c not in cache)
    print(f"{len(valores)} linhas | {len(unicos)} CNPJs únicos | {len(pendentes)} para consultar")

    try:
        for i, cnpj in enumerate(pendentes, 1):
            resultado = consultar(cnpj)
            if not resultado.get("temporario"):
                cache[cnpj] = resultado
            info = montar_colunas(resultado)
            print(f"[{i}/{len(pendentes)}] {cnpj} -> {info.get('telefone_1') or info['status_consulta']}")
            if i % 10 == 0:
                cache_path.write_text(json.dumps(cache, ensure_ascii=False))
            time.sleep(args.intervalo)
    except KeyboardInterrupt:
        print("\nInterrompido. Progresso salvo, rode de novo para continuar.")
    finally:
        cache_path.write_text(json.dumps(cache, ensure_ascii=False))

    fonte.salvar(saida, cache)

    # Resumo por CNPJ único (a mesma empresa aparece em mais de uma aba)
    resumo = [montar_colunas(cache[c]) for c in unicos if c in cache]
    com_tel = sum(bool(i.get("telefone_1") or i.get("telefone_2")) for i in resumo)
    com_cel = sum(bool(i.get("whatsapp_provavel")) for i in resumo)
    print(f"\nPronto! {saida}")
    print(f"  {len(unicos)} empresas | {com_tel} com telefone | {com_cel} com celular (provável WhatsApp)")


if __name__ == "__main__":
    main()
